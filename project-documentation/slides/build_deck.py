#!/usr/bin/env python3
"""Build the presenter deck, the PDF and the PPTX for a week's slides.

Source of truth is the student card page, site/week-NN-slides.html — one card per
slide — plus week-NN-notes.md in this folder, one "## sN" block per card. Nothing is
authored here twice: the cards carry what is projected, the notes carry what is said.

Outputs:
  site/week-NN-deck.html   presenter deck, published beside the card page — arrow keys, deep
                           links (#s5), print CSS. The notes are instructor material: they sit
                           behind the site's instructor view (?instructor) and N toggles them
                           only then. Links stay relative, so the publication gate applies.
  week-NN-deck.pdf         beside this script: the deck printed with the notes (needs Chrome)
  week-NN-deck.pptx        beside this script: the slides with notes in the notes pane

Usage:
  python3 build_deck.py            # both weeks, every output the machine can make
  python3 build_deck.py 2          # one week
  python3 build_deck.py 2 --no-pdf --no-pptx
"""
import argparse
import html
import re
import shutil
import subprocess
import sys
from html.parser import HTMLParser
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
SITE = ROOT / "site"
SITE_URL = "https://datadrivendecisions.github.io/ai-in-business/"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

# ----------------------------------------------------------------- the cards

SLIDE_RE = re.compile(r'<section class="slide([^"]*)" id="(s\d+)">(.*?)</section>', re.S)


def read_cards(week):
    src = (SITE / f"week-{week:02d}-slides.html").read_text(encoding="utf-8")
    title = html.unescape(re.sub(r"<[^>]+>", "", re.search(r'<h1 class="bds-h1">(.*?)</h1>', src, re.S).group(1))).strip()
    lead_m = re.search(r'<p class="lead">(.*?)</p>', src, re.S)
    lead = html.unescape(re.sub(r"<[^>]+>", "", lead_m.group(1))).strip() if lead_m else ""
    cards = []
    for classes, sid, body in SLIDE_RE.findall(src):
        kicker = re.search(r'<div class="slide-kicker">(.*?)</div>', body, re.S).group(1).strip()
        h2 = re.search(r"<h2>(.*?)</h2>", body, re.S).group(1).strip()
        rest = body.split("</h2>", 1)[1].strip()
        cards.append({
            "id": sid,
            "title_slide": "is-title" in classes,
            "kicker": kicker,
            "heading": h2,
            "body": rest,
        })
    return title, lead, cards


def absolutise(fragment):
    """A card links to prd-criteria.html; the deck is opened from a file, so point at the site."""
    return re.sub(r'href="(?!https?://|#)([^"]+)"', lambda m: f'href="{SITE_URL}{m.group(1)}"', fragment)


# ----------------------------------------------------------------- the notes

def read_notes(week):
    path = HERE / f"week-{week:02d}-notes.md"
    if not path.exists():
        return {}
    text = path.read_text(encoding="utf-8")
    notes = {}
    for m in re.finditer(r"^## (s\d+)\s*$(.*?)(?=^## s\d+\s*$|\Z)", text, re.S | re.M):
        notes[m.group(1)] = m.group(2).strip()
    return notes


def md_to_html(md):
    """Paragraphs, '- ' bullets and **bold**; that is all the notes use."""
    out, para, items = [], [], []

    def flush():
        nonlocal para, items
        if para:
            out.append("<p>" + inline(" ".join(para)) + "</p>")
            para = []
        if items:
            out.append("<ul>" + "".join(f"<li>{inline(i)}</li>" for i in items) + "</ul>")
            items = []

    def inline(t):
        t = html.escape(t, quote=False)
        return re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", t)

    for line in md.splitlines():
        if not line.strip():
            flush()
        elif line.startswith("- "):
            if para:
                flush()
            items.append(line[2:].strip())
        else:
            if items:
                flush()
            para.append(line.strip())
    flush()
    return "\n".join(out)


def md_to_text(md):
    lines = []
    for line in md.splitlines():
        line = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        lines.append(("• " + line[2:]) if line.startswith("- ") else line)
    return "\n".join(lines).strip()


# ------------------------------------------------------------ presenter HTML

DECK_CSS = """
:root{--primary:#d4db3e;--supporting:#1a1a1a;--bg-cream:#f5f0e6;--bg-yellow:#fafbf0;--bg-off-white:#f9f8f4;
--text-primary:#111827;--text-secondary:#4b5563;--text-muted:#6b7280;--bg-page:#fff;--border:#e5e7eb;--border-strong:#d1d5db;}
*{box-sizing:border-box}
html,body{margin:0;height:100%;background:#111;color:var(--text-primary);
font-family:-apple-system,BlinkMacSystemFont,'Segoe UI','Inter',Roboto,'Helvetica Neue',Arial,sans-serif}
a{color:inherit;text-decoration:underline;text-decoration-color:var(--primary);text-decoration-thickness:.12em;text-underline-offset:.15em}
.app{display:grid;grid-template-columns:1fr;grid-template-rows:1fr auto;height:100%}
.app.notes-on{grid-template-columns:1fr minmax(18rem,28vw)}
.stagewrap{display:flex;align-items:center;justify-content:center;padding:1.5vmin;min-width:0}
.stage{position:relative;width:min(100%,calc((100vh - 5rem) * 16 / 9));aspect-ratio:16/9;container-type:inline-size;background:var(--bg-page);
box-shadow:0 8px 40px rgba(0,0,0,.5);overflow:hidden}
.slide{position:absolute;inset:0;display:none;flex-direction:column;padding:5cqw 6cqw 6cqw;line-height:1.35}
.slide.is-dense p,.slide.is-dense li{font-size:2.15cqw}.slide.is-dense li{margin-bottom:.9cqw}.slide.is-dense h2{font-size:4cqw;margin-bottom:2cqw}
.slide.is-denser p,.slide.is-denser li{font-size:1.9cqw}.slide.is-denser li{margin-bottom:.7cqw}.slide.is-denser h2{font-size:3.6cqw;margin-bottom:1.6cqw}
.slide.current{display:flex}
.slide-head{display:flex;justify-content:space-between;align-items:baseline;margin-bottom:1.2cqw}
.slide-kicker{font-size:1.5cqw;font-weight:700;letter-spacing:.14em;text-transform:uppercase;color:var(--supporting)}
.slide-kicker::before{content:"";display:inline-block;width:2.2cqw;height:.5cqw;background:var(--primary);border-radius:99px;margin-right:1cqw;vertical-align:middle}
.slide-no{font-size:1.4cqw;color:var(--text-muted);font-variant-numeric:tabular-nums}
.slide h2{font-size:4.6cqw;line-height:1.15;margin:0 0 2.6cqw;font-weight:700}
.slide p,.slide li{font-size:2.5cqw;color:var(--text-primary)}
.slide p{margin:0 0 1.6cqw}
.slide ul{margin:0;padding-left:1.1em}
.slide li{margin-bottom:1.3cqw}
.slide li strong{color:var(--text-primary)}
.slide em{color:var(--text-secondary)}
.slide > :last-child{margin-bottom:0}
.slide.is-title{background:var(--bg-cream);justify-content:center}
.slide.is-title h2{font-size:7cqw}
.slide.is-title p{color:var(--text-secondary);font-size:3cqw}
.cols{display:grid;grid-template-columns:1fr 1fr;gap:4cqw}
.cols h3{font-size:1.5cqw;font-weight:700;letter-spacing:.14em;text-transform:uppercase;color:var(--text-secondary);margin:0 0 1.2cqw}
.slide .roles{display:grid;grid-template-columns:repeat(4,1fr);gap:1.5cqw;margin:1.5cqw 0 2cqw;padding:0;list-style:none}
.slide .roles li{background:var(--bg-cream);border-radius:.8cqw;padding:1.4cqw 1cqw;font-weight:700;text-align:center;margin:0;font-size:2.2cqw}
.tpl{color:var(--text-muted);font-style:italic;border-bottom:.15cqw dashed var(--border-strong)}
li.tpl,p.tpl{display:block;border:.15cqw dashed var(--border-strong);border-radius:.6cqw;padding:1cqw 1.5cqw;background:var(--bg-off-white);list-style:none;margin-left:-1.1em}
.foot{position:absolute;left:6cqw;right:6cqw;bottom:1.8cqw;display:flex;justify-content:space-between;font-size:1.3cqw;color:var(--text-muted)}
.notes{background:#1c1c1c;color:#e5e7eb;padding:1.5rem 1.6rem;overflow:auto;font-size:1.05rem;line-height:1.5;border-left:1px solid #333;display:none}
.notes-on .notes{display:block}
.notes h3{margin:0 0 .8rem;font-size:.8rem;letter-spacing:.14em;text-transform:uppercase;color:#a3a3a3}
.notes p{margin:0 0 .8rem}.notes ul{margin:0 0 .8rem;padding-left:1.2em}.notes li{margin-bottom:.35rem}
.notes .empty{color:#777;font-style:italic}
.notes .iv{font-size:.85rem;color:#a3a3a3;border-left:3px solid #444;padding-left:.7rem}.notes .iv a{color:#ddd}.notes code{font-family:inherit;background:#333;padding:0 .3em;border-radius:3px}
.bar{grid-column:1/-1;display:flex;align-items:center;gap:1rem;padding:.5rem 1rem;background:#000;color:#bbb;font-size:.85rem}
.bar button{background:#222;color:#eee;border:1px solid #444;border-radius:6px;padding:.3rem .7rem;cursor:pointer;font:inherit}
.bar button:hover{background:#333}
.bar .counter{font-variant-numeric:tabular-nums;min-width:4.5rem;text-align:center}
.bar .title{margin-left:auto;color:#888;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.bar kbd{background:#222;border:1px solid #444;border-radius:4px;padding:0 .35em;font-family:inherit}
.printonly{display:none}
@media print{
  @page{size:A4 landscape;margin:12mm}
  html,body{background:#fff;height:auto}
  .app,.bar,.notes,.stagewrap{display:none!important}
  .printonly{display:block}
  .page{break-after:page;display:grid;grid-template-rows:auto 1fr;gap:4mm;min-height:calc(210mm - 24mm)}
  .page:last-child{break-after:auto}
  .page .stage{width:190mm;aspect-ratio:16/9;box-shadow:none;border:.3mm solid #999;margin:0 auto;container-type:inline-size;position:relative;background:#fff;overflow:hidden}
  .page .slide{display:flex}
  .page .pnotes{font-size:10pt;line-height:1.4;color:#111;column-count:2;column-gap:8mm}
  .page .pnotes h3{font-size:8pt;letter-spacing:.14em;text-transform:uppercase;color:#666;margin:0 0 2mm;column-span:all}
  .page .pnotes p{margin:0 0 2.5mm}.page .pnotes ul{margin:0 0 2.5mm;padding-left:1.2em}
}
"""

def instructor_js():
    """The site's instructor-view toggle, copied verbatim from the foot of index.html at build time."""
    src = (SITE / "index.html").read_text(encoding="utf-8")
    return src[src.rfind("<script>"):src.rfind("</script>") + len("</script>")]


DECK_JS = """
(function(){
  var slides=[].slice.call(document.querySelectorAll('#app .stage > .slide'));
  var app=document.getElementById('app'), counter=document.getElementById('counter');
  var noteBoxes=[].slice.call(document.querySelectorAll('.notes .note'));
  var i=0;
  function go(n){
    i=Math.max(0,Math.min(slides.length-1,n));
    slides.forEach(function(s,k){s.classList.toggle('current',k===i)});
    noteBoxes.forEach(function(b,k){b.hidden=(k!==i)});
    counter.textContent=(i+1)+' / '+slides.length;
    if(location.hash!=='#'+slides[i].id) history.replaceState(null,'','#'+slides[i].id);
  }
  function fromHash(){var k=slides.findIndex(function(s){return '#'+s.id===location.hash});go(k<0?0:k)}
  var notesEl=document.querySelector('.notes');
  function notesAllowed(){return notesEl&&!notesEl.hidden}
  function toggleNotes(){if(!notesAllowed())return;app.classList.toggle('notes-on');try{localStorage.setItem('deck-notes',app.classList.contains('notes-on')?'1':'0')}catch(e){}}
  document.addEventListener('keydown',function(e){
    if(e.metaKey||e.ctrlKey||e.altKey) return;
    switch(e.key){
      case 'ArrowRight': case 'ArrowDown': case 'PageDown': case ' ': case 'Enter': go(i+1); e.preventDefault(); break;
      case 'ArrowLeft': case 'ArrowUp': case 'PageUp': case 'Backspace': go(i-1); e.preventDefault(); break;
      case 'Home': go(0); break; case 'End': go(slides.length-1); break;
      case 'n': case 'N': toggleNotes(); break;
      case 'f': case 'F': if(document.fullscreenElement) document.exitFullscreen(); else document.documentElement.requestFullscreen(); break;
      case 'p': case 'P': window.print(); break;
    }
  });
  document.getElementById('prev').onclick=function(){go(i-1)};
  document.getElementById('next').onclick=function(){go(i+1)};
  document.getElementById('notesbtn').onclick=toggleNotes;
  document.getElementById('printbtn').onclick=function(){window.print()};
  document.querySelector('#app .stage').addEventListener('click',function(e){
    if(e.target.closest('a')) return;
    var r=this.getBoundingClientRect(); go(e.clientX-r.left>r.width/2?i+1:i-1);
  });
  window.addEventListener('hashchange',fromHash);
  var want=new URLSearchParams(location.search).get('notes');
  try{ if(notesAllowed()&&(want==='1'||(want===null&&localStorage.getItem('deck-notes')==='1'))) app.classList.add('notes-on'); }catch(e){}
  fromHash();
})();
"""


def slide_html(card, n, total, course_line):
    cls = "slide is-title" if card["title_slide"] else "slide"
    words = len(html.unescape(re.sub(r"<[^>]+>", " ", card["body"])).split())
    cls += " is-denser" if words > 85 else (" is-dense" if words > 55 else "")
    return (
        f'<section class="{cls}" id="{card["id"]}">'
        f'<div class="slide-head"><div class="slide-kicker">{card["kicker"]}</div><span class="slide-no">{n} / {total}</span></div>'
        f'<h2>{card["heading"]}</h2>{card["body"]}'
        f'<div class="foot"><span>{course_line}</span><span>{n} / {total}</span></div>'
        f"</section>"
    )


def build_html(week, title, lead, cards, notes):
    total = len(cards)
    course_line = f"AI in Business · week {week}"
    slides = "\n".join(slide_html(c, k + 1, total, course_line) for k, c in enumerate(cards))
    note_boxes = "\n".join(
        f'<div class="note" id="note-{c["id"]}"{"" if k == 0 else " hidden"}>'
        + (md_to_html(notes[c["id"]]) if notes.get(c["id"]) else '<p class="empty">No notes for this slide.</p>')
        + "</div>"
        for k, c in enumerate(cards)
    )
    # The print view repeats every slide inside its own page, notes underneath.
    pages = "\n".join(
        f'<div class="page"><div class="stage">{slide_html(c, k + 1, total, course_line).replace("class=\"slide", "class=\"slide current", 1)}</div>'
        f'<div class="pnotes" data-audience="instructor" hidden><h3>Notes · {k + 1} / {total}</h3>'
        + (md_to_html(notes[c["id"]]) if notes.get(c["id"]) else "<p><em>No notes.</em></p>")
        + "</div></div>"
        for k, c in enumerate(cards)
    )
    INSTRUCTOR_JS = instructor_js()
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8" />
<meta name="viewport" content="width=device-width, initial-scale=1.0" />
<title>{html.escape(title)} — week {week} presenter deck</title>
<meta name="robots" content="noindex" />
<style>{DECK_CSS}</style>
</head>
<body>
<div class="app" id="app">
  <div class="stagewrap"><div class="stage">
{slides}
  </div></div>
  <aside class="notes" aria-label="Instructor notes" data-audience="instructor" hidden>
    <p class="iv"><strong>Instructor view.</strong> The notes show because the address carried <code>?instructor</code>. <a href="week-{week:02d}-deck.html?student">Leave the instructor view</a>.</p>
    <h3>Instructor notes</h3>
{note_boxes}
  </aside>
  <div class="bar">
    <button id="prev" type="button" aria-label="Previous slide">◀</button>
    <span class="counter" id="counter">1 / {total}</span>
    <button id="next" type="button" aria-label="Next slide">▶</button>
    <button id="notesbtn" type="button" data-audience="instructor" hidden>Notes <kbd>N</kbd></button>
    <button id="printbtn" type="button">Print <kbd>P</kbd></button>
    <span><kbd>←</kbd> <kbd>→</kbd> or click · <kbd>F</kbd> full screen</span>
    <span class="title"><a href="week-{week:02d}-slides.html">{html.escape(title)} · week {week} · one card per slide</a> · <a href="index.html">course</a></span>
  </div>
</div>
<div class="printonly">
{pages}
</div>
{INSTRUCTOR_JS}
<script>{DECK_JS}</script>
</body>
</html>
"""


# ------------------------------------------------------------------- the PPTX

class Blocks(HTMLParser):
    """Card body HTML -> [(column, kind, runs)], runs = [(text, bold, italic, href, muted)]."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks, self.col, self.stack, self.cur = [], 0, [], None
        self.in_cols, self.col_depth, self.roles = False, 0, False

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        cls = a.get("class", "")
        if tag == "div" and "cols" in cls:
            self.in_cols, self.col = True, -1
        elif tag == "div" and self.in_cols:
            self.col += 1
        elif tag == "ul" and "roles" in cls:
            self.roles = True
            self.cur = [self.col, "roles", []]
        elif tag in ("p", "li", "h3") and not self.roles:
            self.cur = [max(self.col, 0), tag, []]
            if "tpl" in cls:
                self.stack.append(("tpl", None))
        elif tag in ("strong", "em", "a", "span"):
            self.stack.append((tag if tag != "span" or "tpl" not in cls else "tpl", a.get("href")))
        elif tag == "br" and self.cur:
            self.cur[2].append(("\n", False, False, None, False))

    def handle_endtag(self, tag):
        if tag in ("strong", "em", "a", "span") and self.stack:
            self.stack.pop()
        elif tag in ("p", "li", "h3") and self.cur and not self.roles:
            if self.stack and self.stack[-1][0] == "tpl" and tag != "span":
                self.stack.pop()
            self.blocks.append(tuple(self.cur)); self.cur = None
        elif tag == "li" and self.roles and self.cur:
            self.cur[2].append((" · ", True, False, None, False))
        elif tag == "ul" and self.roles:
            runs = self.cur[2][:-1] if self.cur[2] else []
            self.blocks.append((self.cur[0], "roles", runs)); self.cur, self.roles = None, False
        elif tag == "div" and self.in_cols and self.col >= 0 and tag == "div":
            pass

    def handle_data(self, data):
        if self.cur is None:
            if data.strip():
                self.cur = [max(self.col, 0), "p", []]
            else:
                return
        text = re.sub(r"\s+", " ", data)
        bold = any(t == "strong" for t, _ in self.stack) or self.roles
        italic = any(t == "em" for t, _ in self.stack)
        href = next((h for t, h in self.stack if t == "a" and h), None)
        muted = any(t == "tpl" for t, _ in self.stack)
        self.cur[2].append((text, bold, italic, href, muted))


def body_blocks(fragment):
    p = Blocks(); p.feed(absolutise(fragment)); p.close()
    if p.cur and p.cur[2]:
        p.blocks.append(tuple(p.cur))
    return p.blocks


def build_pptx(week, title, cards, notes, out):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    INK, SEC, MUTED = RGBColor(0x11, 0x18, 0x27), RGBColor(0x4B, 0x55, 0x63), RGBColor(0x6B, 0x72, 0x80)
    LIME, CREAM, DARK = RGBColor(0xD4, 0xDB, 0x3E), RGBColor(0xF5, 0xF0, 0xE6), RGBColor(0x1A, 0x1A, 0x1A)
    W, H = Inches(13.333), Inches(7.5)
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    total = len(cards)

    def textbox(slide, x, y, w, h):
        tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
        return tf

    def add_runs(par, runs, size, color=INK):
        for text, bold, italic, href, muted in runs:
            if not text:
                continue
            r = par.add_run(); r.text = text
            r.font.size = Pt(size); r.font.bold = bold or None; r.font.italic = italic or muted or None
            r.font.color.rgb = MUTED if muted else (SEC if italic else color)
            if href:
                r.hyperlink.address = href

    for k, card in enumerate(cards):
        s = prs.slides.add_slide(blank)
        if card["title_slide"]:
            bg = s.background.fill; bg.solid(); bg.fore_color.rgb = CREAM
        # brand mark + kicker
        mark = s.shapes.add_shape(1, Inches(0.8), Inches(0.62), Inches(0.35), Inches(0.08))
        mark.fill.solid(); mark.fill.fore_color.rgb = LIME; mark.line.fill.background()
        tf = textbox(s, Inches(1.25), Inches(0.45), Inches(8), Inches(0.4))
        p = tf.paragraphs[0]; r = p.add_run(); r.text = html.unescape(re.sub(r"<[^>]+>", "", card["kicker"])).upper()
        r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = DARK
        tf = textbox(s, Inches(10.5), Inches(0.45), Inches(2.0), Inches(0.4))
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT; r = p.add_run(); r.text = f"{k + 1} / {total}"
        r.font.size = Pt(11); r.font.color.rgb = MUTED
        # heading
        heading = html.unescape(re.sub(r"<[^>]+>", "", card["heading"]))
        tf = textbox(s, Inches(0.8), Inches(1.0) if not card["title_slide"] else Inches(2.4), Inches(11.7), Inches(1.3))
        p = tf.paragraphs[0]; r = p.add_run(); r.text = heading
        r.font.size = Pt(44 if card["title_slide"] else 32); r.font.bold = True; r.font.color.rgb = INK
        # body
        blocks = body_blocks(card["body"])
        ncols = max((b[0] for b in blocks), default=0) + 1
        top = Inches(3.8) if card["title_slide"] else Inches(2.3)
        colw = (Inches(11.7) - Inches(0.5) * (ncols - 1)) / ncols
        frames = [textbox(s, Inches(0.8) + (colw + Inches(0.5)) * c, top, colw, H - top - Inches(0.9)) for c in range(ncols)]
        first = [True] * ncols
        for col, kind, runs in blocks:
            tf = frames[col]
            p = tf.paragraphs[0] if first[col] else tf.add_paragraph(); first[col] = False
            p.space_after = Pt(8)
            if kind == "h3":
                add_runs(p, [(t.upper(), True, False, h, m) for t, b, i, h, m in runs], 11, SEC)
            elif kind == "li":
                add_runs(p, [("•  ", False, False, None, False)] + list(runs), 18)
            elif kind == "roles":
                add_runs(p, runs, 20)
            else:
                add_runs(p, runs, 22 if card["title_slide"] else 18, SEC if card["title_slide"] else INK)
        # footer
        tf = textbox(s, Inches(0.8), H - Inches(0.6), Inches(8), Inches(0.35))
        r = tf.paragraphs[0].add_run(); r.text = f"AI in Business · week {week}"; r.font.size = Pt(9); r.font.color.rgb = MUTED
        # notes
        s.notes_slide.notes_text_frame.text = md_to_text(notes.get(card["id"], ""))
    prs.save(out)


# -------------------------------------------------------------------- the PDF

def build_pdf(deck_html, out):
    if not Path(CHROME).exists():
        return "skipped: Google Chrome not found"
    r = subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={out}", f"file://{deck_html}?instructor"], capture_output=True, text=True, timeout=120)
    return "ok" if out.exists() else f"failed: {r.stderr[-300:]}"


# ------------------------------------------------------------------------ main

def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("weeks", nargs="*", type=int, default=[1, 2])
    ap.add_argument("--no-pdf", action="store_true"); ap.add_argument("--no-pptx", action="store_true")
    a = ap.parse_args(argv)
    for week in a.weeks:
        title, lead, cards = read_cards(week)
        notes = read_notes(week)
        missing = [c["id"] for c in cards if c["id"] not in notes]
        deck = SITE / f"week-{week:02d}-deck.html"
        deck.write_text(build_html(week, title, lead, cards, notes), encoding="utf-8")
        print(f"week {week}: {len(cards)} slides, notes for {len(cards) - len(missing)}"
              + (f" (missing: {', '.join(missing)})" if missing else "") + f" → {deck.name}")
        if not a.no_pdf:
            print(f"  pdf  → {build_pdf(deck, HERE / f'week-{week:02d}-deck.pdf')}")
        if not a.no_pptx:
            try:
                build_pptx(week, title, cards, notes, HERE / f"week-{week:02d}-deck.pptx"); print("  pptx → ok")
            except ImportError:
                print("  pptx → skipped: python-pptx not installed (pip install python-pptx)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
